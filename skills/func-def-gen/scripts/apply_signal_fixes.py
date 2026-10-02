# -*- coding: utf-8 -*-
"""
func-def-gen · 改动应用与 bot 批注标记脚本（合并单元格安全版）

职责（对应 SKILL.md 工作流第 6.5 步）：
  1. 信号表报文/CANID 校正 —— 整行内容匹配（任意格==旧报文 且 任意格以信号名开头），
     替换行内所有旧报文/旧 CANID 格。绝不按固定列索引定位（合并单元格行会列位错位漏改）。
  2. 段落正文改写 —— run-aware 替换，可用 chapter 限定 Heading2 章节作用域
     （防止类似"PWC网络节点丢失"误伤手机无线充电章的事故）。
  3. 单元格值改写 —— 按行信号名前缀 + 表头列名定位。
  4. 每处改动自动加 Word 批注（默认作者 bot），写明改动内容/前后值/依据。
  5. 改完自动做终态全量核对：任何含目标信号前缀的行若仍含旧报文 → 报错退出。
     （教训：不能只信单次改写计数，必须全量复查。）

用法:
  python apply_signal_fixes.py <config.json> [--dry-run]
config 字段见 SKILL.md「config 字段说明」之 signal_fixes / para_fixes / cell_fixes。
"""
import sys, json, re
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table


def _hexint(s):
    """从文本提取十六进制数（用于CANID列识别）。"""
    if s is None:
        return None
    m = re.search(r'0[xX]([0-9A-Fa-f]+)', str(s))
    if not m:
        return None
    try:
        return int(m.group(1), 16)
    except ValueError:
        return None

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

def get_bot_comment_ids(doc, author):
    """已存在的、由本脚本作者挂的批注 id 集合（幂等防重挂）。"""
    ids = set()
    try:
        for part in doc.part.package.iter_parts():
            if str(part.partname) == '/word/comments.xml':
                root = etree.fromstring(part.blob)
                for c in root.findall(f'{W}comment'):
                    if c.get(f'{W}author') == author:
                        ids.add(c.get(f'{W}id'))
                break
    except Exception:
        pass
    return ids

def row_has_bot_comment(row, bot_ids):
    """该行是否已挂过 bot 批注（锚点 commentRangeStart 的 id 命中）。"""
    if not bot_ids:
        return False
    for c in dedupe_cells(row.cells):
        for p in c.paragraphs:
            for cs in p._p.findall(f'{W}commentRangeStart'):
                if cs.get(f'{W}id') in bot_ids:
                    return True
    return False

def dedupe_cells(row_or_cells):
    """合并单元格去重：同一 _tc 连续重复只保留一次，使数据行列位与表头对齐。"""
    seen = []
    for c in row_or_cells:
        if not seen or c._tc is not seen[-1]._tc:
            seen.append(c)
    return seen

def set_cell_text(cell, text):
    """run-aware 写格：保留首个 run 格式，清空其余。"""
    first = False
    for p in cell.paragraphs:
        if not first and p.runs:
            p.runs[0].text = text
            for r in p.runs[1:]:
                r.text = ''
            first = True
        else:
            for r in p.runs:
                r.text = ''
    if not first:
        cell.paragraphs[0].add_run(text)

def replace_in_para(p, old, new):
    """run-aware 段内子串替换。命中返回 True。"""
    full = ''.join(r.text for r in p.runs)
    if old not in full:
        return False
    full = full.replace(old, new)
    if p.runs:
        p.runs[0].text = full
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(full)
    return True

def sig_match(text, sigp, kw=''):
    """信号名完整词边界包含匹配：sigp 后不能紧跟字母/数字/下划线，
    防止 SupEduFunSwt 误匹配 SupEduFunSwtEnaFlg。"""
    return re.search(re.escape(sigp) + r'(?![A-Za-z0-9_])', text) is not None \
            and kw in text

def main():
    cfg = json.load(open(sys.argv[1], encoding='utf-8'))
    dry = '--dry-run' in sys.argv
    doc = Document(cfg['src'])
    AUTHOR = cfg.get('comment_author', 'bot')
    fixes = cfg.get('signal_fixes', [])
    pfixes = cfg.get('para_fixes', [])
    cfixes = cfg.get('cell_fixes', [])
    notes = cfg.get('signal_notes', [])   # 纯提示批注（不改值），如"矩阵无对应信号"
    log = {'sig': [], 'para': [], 'cell': [], 'note': [], 'comments': 0}
    bot_ids = get_bot_comment_ids(doc, AUTHOR)

    def add_comment(runs, text):
        if dry or not runs:
            return
        doc.add_comment(runs=runs, text=text, author=AUTHOR, initials=AUTHOR)
        log['comments'] += 1

    # ---- 0) 幂等与终态核对共用的按值判据（与是否已有批注无关） ----
    def rule_pending(f, cells, texts):
        """判断一条 signal_fix 规则在该行上是否仍有未落地的改动。
        幂等跳过与终态 residual 共用此判据，防止两处口径打架。"""
        oldm, newm = f['old_msg'], f['new_msg']
        oldc, newc = f.get('old_canid', ''), f.get('new_canid', '')
        olds, news = f.get('old_sig'), f.get('new_sig')
        why = []
        if oldm != newm and (oldm in texts or (oldc and oldc != newc and oldc in texts)):
            why.append('旧报文/' + oldm)
        if oldm == newm and oldc and newc and oldc != newc and oldc in texts:
            why.append('旧CANID/' + oldc)
        if olds and news and olds != news:
            pat = r'(?<![A-Za-z0-9_])' + re.escape(olds) + r'(?![A-Za-z0-9_])'
            sig_cell = next((c for c in cells
                             if sig_match(c.text.strip(), f['signal_prefix'], f.get('keyword', ''))), None)
            if sig_cell is not None and re.search(pat, sig_cell.text):
                why.append('旧信号名/' + olds)
        return (bool(why), why)

    # ---- 1) 信号表整行校正 ----
    fixed_rows = set()   # 本运行已做过替换的行，1.5 不再挂提示批注
    for t in doc.tables:
        for row in t.rows:
            cells = dedupe_cells(row.cells)
            texts = [c.text.strip() for c in cells]
            for f in fixes:
                sigp, kw = f['signal_prefix'], f.get('keyword', '')
                if not any(sig_match(x, sigp, kw) for x in texts):
                    continue
                # 幂等（按值，与批注无关）：历史中断运行可能"批注已挂、改动未落"，
                # 按批注跳过会把该行永久卡死；已落地的行（含已挂批注的）在此自然跳过。
                if not rule_pending(f, cells, texts)[0]:
                    continue
                oldm, oldc = f['old_msg'], f.get('old_canid', '')
                newm, newc = f['new_msg'], f.get('new_canid', '')
                # CANID 列一并更新：旧CANID格(值==old_canid)→new_canid，
                # 且报文变化时该报文同名的旧CANID必须跟着变（合并单元格行尤其如此）。
                hit = [c for c in cells if c.text.strip() in (oldm, oldc)]
                if oldm != newm and f.get('new_canid'):
                    ocv = _hexint(oldc) if oldc else None
                    if ocv is None:
                        # old_canid 未从报文名推断时：取报文名里 0x 后面的部分
                        mm = re.search(r'0[xX]([0-9A-Fa-f]+)', oldm)
                        ocv = int(mm.group(1), 16) if mm else None
                    ncv = _hexint(f['new_canid'])
                    if ocv is not None and ncv is not None:
                        for c in cells:
                            if c in hit:
                                continue
                            cv = _hexint(c.text)
                            if cv is not None and cv == ocv:
                                hit.append(c)
                sig_cell = next((c for c in cells if sig_match(c.text.strip(), sigp)), None)
                need_sig_rename = bool(f.get('old_sig')) and sig_cell is not None \
                        and f['old_sig'] in sig_cell.text and f.get('new_sig') not in sig_cell.text
                if not hit and not need_sig_rename:
                    continue
                if dry:
                    log['sig'].append(f"[dry] {sigp}: {oldm}->{newm} {f.get('old_sig','')}->{f.get('new_sig','')}")
                    continue
                runs = []
                for c in hit + ([sig_cell] if need_sig_rename else []):
                    for p in c.paragraphs:
                        runs.extend(p.runs)
                for c in hit:
                    tx = c.text.strip()
                    set_cell_text(c, newm if tx == oldm else newc)
                if need_sig_rename:
                    for p in sig_cell.paragraphs:
                        replace_in_para(p, f['old_sig'], f['new_sig'])
                add_comment(runs, f.get('comment', f"报文更正：{oldm}→{newm}"))
                log['sig'].append(f"{sigp}: {oldm}->{newm} {f.get('old_sig','')}->{f.get('new_sig','')}")
                fixed_rows.add(row._tr)

    # ---- 1.5) 信号行纯提示批注（不改值，如"矩阵无对应"） ----
    for t in doc.tables:
        for row in t.rows:
            if row._tr in fixed_rows or row_has_bot_comment(row, bot_ids):
                continue    # 已替换过/已挂过 bot 批注的行不再挂提示
            cells = dedupe_cells(row.cells)
            for f in notes:
                sigp, kw = f['signal_prefix'], f.get('keyword', '')
                sig_cell = next((c for c in cells if sig_match(c.text.strip(), sigp, kw)), None)
                if sig_cell is None:
                    continue
                if dry:
                    log['note'].append(f"[dry] {sigp}: {f['note'][:40]}")
                    break
                runs = [r for p in sig_cell.paragraphs for r in p.runs]
                add_comment(runs, f['note'])
                log['note'].append(f"{sigp}: {f['note'][:40]}")
                break   # 一行只挂一条提示批注

    # ---- 2) 段落改写（支持 chapter 作用域） ----
    cur_ch = None
    for ch in doc.element.body.iterchildren():
        if ch.tag == qn('w:p'):
            p = Paragraph(ch, doc)
            st = p.style.name if p.style else ''
            if st == 'Heading 2':
                cur_ch = p.text.strip()
                continue
            for f in pfixes:
                if f.get('chapter') and cur_ch != f['chapter']:
                    continue
                sc = f.get('scope_contains')
                if sc and sc not in p.text:
                    continue
                if replace_in_para(p, f['find'], f['replace']):
                    add_comment(p.runs, f.get('comment', f"改写：{f['find']}→{f['replace']}"))
                    log['para'].append(f"[{cur_ch}] {f['find']}→{f['replace']}")

    # ---- 3) 单元格值改写 ----
    for t in doc.tables:
        hdr = [c.text.strip() for c in dedupe_cells(t.rows[0].cells)]
        for f in cfixes:
            col = f.get('column', '信号值描述')
            if f.get('match_signal_prefix', '信号名称') not in hdr or col not in hdr:
                continue
            si = hdr.index(f.get('match_signal_prefix', '信号名称'))
            vi = hdr.index(col)
            for row in t.rows:
                cs = dedupe_cells(row.cells)
                if si < len(cs) and cs[si].text.strip().startswith(f['row_signal_prefix']) \
                        and f['find'] in cs[vi].text:
                    runs = [r for p in cs[vi].paragraphs for r in p.runs]
                    if not dry:
                        for p in cs[vi].paragraphs:
                            replace_in_para(p, f['find'], f['replace'])
                    add_comment(runs, f.get('comment', f"{f['find']}→{f['replace']}"))
                    log['cell'].append(f"{f['row_signal_prefix']}: {f['find']}→{f['replace']}")

    # ---- 4) 终态全量核对（防漏改，严格全表、与批注无关） ----
    residual = []
    for t in doc.tables:
        for row in t.rows:
            texts = [c.text.strip() for c in dedupe_cells(row.cells)]
            for f in fixes:
                sigp, kw = f['signal_prefix'], f.get('keyword', '')
                if not any(sig_match(x, sigp, kw) for x in texts):
                    continue
                pending, why = rule_pending(f, dedupe_cells(row.cells), texts)
                if pending:
                    residual.append(f"{sigp} 行仍有未落地改动（{'、'.join(why)}）")
    if residual and not dry:
        print("!! 终态核对失败，存在漏改：")
        for r in residual:
            print("  -", r)
        doc.save(cfg['out'])
        sys.exit(2)

    if not dry:
        doc.save(cfg['out'])
    print(f"信号行校正: {len(log['sig'])} | 段落改写: {len(log['para'])} | 单元格改写: {len(log['cell'])} | 提示批注: {len(log['note'])} | 批注合计: {log['comments']} (作者={AUTHOR}){' [DRY-RUN]' if dry else ''}")
    for k in ('sig', 'para', 'cell', 'note'):
        for l in log[k]:
            print(f"  {l}")

if __name__ == '__main__':
    main()
