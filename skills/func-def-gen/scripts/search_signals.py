# -*- coding: utf-8 -*-
"""Search an IBUS/CANFD signal matrix xlsx for relevant signals.
Usage: python search_signals.py <matrix.xlsx> <kw1,kw2,...> <out.txt>
Output: msg=... | sig=... | cn=... | en=... | desc=...
Iterates the 'Matrix' sheet, tracks current message name, dedup by (msg,sig).
"""
import sys, re
from openpyxl import load_workbook

def main():
    if len(sys.argv) < 4:
        print("usage: search_signals.py <matrix.xlsx> <kw1,kw2> <out.txt>")
        sys.exit(1)
    path, kws, outp = sys.argv[1], sys.argv[2], sys.argv[3]
    keywords = [k.strip() for k in kws.split(',') if k.strip()]
    pat = re.compile("(" + "|".join(keywords) + ")", re.I)
    wb = load_workbook(path, data_only=True)
    ws = wb["Matrix"]
    hdr = {}
    for cell in ws[1]:
        if cell.value:
            hdr[cell.column] = str(cell.value).replace('\n', ' ')
    def ci(p):
        return next((c for c, h in hdr.items() if p in h), None)
    CI = {'MSG': ci("Msg Name"), 'SIG': ci("Signal Name"), 'CN': ci("Chinese"),
          'EN': ci("English"), 'DESC': ci("Signal Value Description")}
    cur = None
    seen = set()
    lines = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        def g(k):
            c = CI[k]
            return row[c-1] if c and len(row) >= c else None
        m = g('MSG')
        if m and str(m).strip():
            cur = str(m).strip()
        sig = g('SIG')
        if not sig or str(sig).strip() == "":
            continue
        sig = str(sig).strip()
        blob = (cur + " " + sig + " " + (str(g('CN')) if g('CN') else "") + " " + (str(g('EN')) if g('EN') else ""))
        if pat.search(blob):
            key = (cur, sig)
            if key in seen:
                continue
            seen.add(key)
            desc = str(g('DESC'))[:120] if g('DESC') else ""
            lines.append(f"msg={cur}|sig={sig}|cn={g('CN')}|en={g('EN')}|desc={desc}")
    with open(outp, 'w', encoding='utf-8') as f:
        f.write(f"# matched {len(lines)} signals for keywords: {keywords}\n")
        for l in lines:
            f.write(l + "\n")
    print(f"matched {len(lines)} signals -> {outp}")

if __name__ == '__main__':
    main()
