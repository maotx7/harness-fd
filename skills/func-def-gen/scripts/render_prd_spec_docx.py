"""Render a validated PRD function-definition spec into a standalone DOCX.

This renderer deliberately does not load a project template.  The semantic
contract is the source of truth; Word is only the readable delivery format.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SCRIPT_DIR = Path(__file__).resolve().parent


def _load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法加载模块: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


_contract = _load_module("prd_spec_contract", SCRIPT_DIR / "workflow_contract.py")
_quality = _load_module("prd_spec_quality", SCRIPT_DIR / "prd_quality_verify.py")


def _strings(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value.strip()] if value.strip() else []
    if isinstance(value, list):
        result: list[str] = []
        for item in value:
            result.extend(_strings(item))
        return result
    if isinstance(value, dict):
        result = []
        for key, item in value.items():
            values = _strings(item)
            result.extend([f"{key}: {text}" for text in values])
        return result
    return [str(value).strip()] if value is not None and str(value).strip() else []


def _safe_filename(title: str) -> str:
    value = re.sub(r"[\\/:*?\"<>|\r\n]+", "_", title).strip(" ._")
    return value or "功能定义"


def _set_font(run, *, size: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _paragraph(document: Document, text: str, *, style: str | None = None, bold: bool = False):
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(str(text))
    _set_font(run, bold=bold)
    return paragraph


def _add_value(document: Document, value: Any) -> None:
    values = _strings(value)
    if not values:
        _paragraph(document, "（未提供）")
        return
    for item in values:
        if "\n" in item:
            for line in item.splitlines():
                if line.strip():
                    _paragraph(document, line.strip())
        else:
            _paragraph(document, item)


def _style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 20), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        if name in styles:
            styles[name].font.name = "黑体"
            styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            styles[name].font.size = Pt(size)


def _heading(document: Document, text: str, level: int, number: str | None = None) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    label = f"{number} {text}" if number else text
    run = paragraph.add_run(label)
    _set_font(run, size={1: 16, 2: 14, 3: 12}.get(level, 11), bold=True)


def _table(document: Document, headers: list[str], rows: list[list[Any]]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = str(header)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            _set_font(run, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = "\n".join(_strings(value))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_font(run)
    return table


def _render_scenario(document: Document, scenario: dict[str, Any], number: str) -> None:
    _heading(document, f"{scenario['id']} {scenario['title']}", 3, number=number)
    _paragraph(document, f"场景目的：{scenario['purpose']}")
    rows = [
        ["前置条件", scenario.get("preconditions")],
        ["使能条件", scenario.get("enable_conditions")],
        ["触发条件", scenario.get("triggers")],
        ["执行步骤", scenario.get("steps")],
        ["输出结果", scenario.get("outputs")],
        ["完成条件", scenario.get("completion_conditions")],
        ["退出条件", scenario.get("exit_conditions")],
        ["异常处理", scenario.get("exceptions")],
        ["恢复条件", scenario.get("recovery")],
    ]
    _table(document, ["场景要素", "定义"], rows)


def _render_signal_table(document: Document, signals: dict[str, dict[str, Any]], keys: list[str]) -> None:
    selected = [signals[key] for key in keys if key in signals]
    if not selected:
        return
    names = [str(item.get("signal_name") or item.get("suggested_name") or "") for item in selected]
    _paragraph(document, f"本节涉及信号：{'、'.join(names)}")
    _table(
        document,
        ["语义键", "语义名称", "信号名称", "状态", "方向", "生产者", "消费者", "用途"],
        [
            [
                item.get("semantic_key"), item.get("semantic_name"),
                item.get("signal_name") or item.get("suggested_name"), item.get("status"),
                item.get("direction"), item.get("producer", ""), item.get("consumer", ""), item.get("purpose"),
            ]
            for item in selected
        ],
    )


def _render_traceability(document: Document, normalized: dict[str, Any], section_name: str) -> None:
    """Render a chapter-7 table below its already-numbered section heading."""
    if section_name == "需求追溯":
        _table(
            document,
            ["需求ID", "需求内容", "来源", "映射子功能"],
            [
                [
                    item["id"], item["text"], item.get("source", ""),
                    ", ".join(sf_id for sf_id, sf in normalized["subfunctions"].items() if item["id"] in sf["requirement_ids"]),
                ]
                for item in normalized["requirements"].values()
            ],
        )
    elif section_name == "信号追溯":
        _table(
            document,
            ["语义键", "信号名称", "状态", "引用场景"],
            [
                [
                    item["semantic_key"], item.get("signal_name") or item.get("suggested_name"), item["status"],
                    ", ".join(sid for sid, scenario in normalized["scenarios"].items() if item["semantic_key"] in scenario["signal_keys"]),
                ]
                for item in normalized["signals"].values()
            ],
        )
    elif section_name == "未决事项":
        unresolved = [item for item in normalized["signals"].values() if item["status"] != "confirmed"]
        if unresolved:
            _table(
                document,
                ["事项", "状态", "说明"],
                [[item["semantic_name"], item["status"], item.get("note", "需补充信号矩阵或确认接口属性")] for item in unresolved],
            )
        else:
            _paragraph(document, "无未决事项。")


def _ordered_section_names(chapter_id: str, sections: dict[str, Any]) -> list[str]:
    """Return contract order, followed by any extra visible sections."""
    chapter_def = next(item for item in _contract.PRD_SPEC_CHAPTERS if item["id"] == chapter_id)
    required = list(chapter_def["sections"])
    extras = [
        name for name in sections
        if name != "signal_keys" and name not in required
    ]
    return required + extras


def render_prd_spec_document(spec_path: str | Path, output_dir: str | Path, output_name: str | None = None) -> Path:
    spec_file = Path(spec_path).expanduser().resolve()
    payload = json.loads(spec_file.read_text(encoding="utf-8"))
    normalized = _contract.normalize_prd_spec_contract(payload)
    report = _quality.review_prd_spec(normalized)
    output_root = Path(output_dir).expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    quality_path = output_root / "quality_report.md"
    _quality._write_report(quality_path, report)
    if not report["ok"]:
        raise _contract.WorkflowContractError(f"PRD功能定义质量门禁未通过，阻断问题: {report['error_count']}")

    title = str(payload.get("title") or "汽车功能定义").strip()
    output = output_root / (output_name or f"{_safe_filename(title)}.docx")
    document = Document()
    _style_document(document)
    document.core_properties.title = title
    cover = document.add_paragraph(style="Title")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover.add_run(title)
    _set_font(run, size=20, bold=True)
    _paragraph(document, f"项目：{payload.get('project') or '未指定'}", bold=True)
    _paragraph(document, f"文档状态：{'草稿' if normalized['mode'] == 'draft' else '正式版'}")
    _paragraph(document, "本文档依据 PRD 结构化分析结果生成，章节内容不依赖既有 Word 模板。")

    chapter_titles = {str(item["id"]): item["title"] for item in _contract.PRD_SPEC_CHAPTERS}
    for chapter_id in sorted(normalized["chapters"], key=lambda value: int(value)):
        chapter = normalized["chapters"][chapter_id]
        _heading(document, chapter.get("title") or chapter_titles[chapter_id], 1, number=chapter_id)
        section_names = _ordered_section_names(chapter_id, chapter["sections"])
        for section_index, section_name in enumerate(section_names, start=1):
            value = chapter["sections"].get(section_name)
            section_number = f"{chapter_id}.{section_index}"
            _heading(document, section_name, 2, number=section_number)
            _add_value(document, value)
            if chapter_id == "5" and section_name == "信号定义和信号状态":
                _render_signal_table(document, normalized["signals"], chapter.get("signal_keys", []))
            if chapter_id == "6" and section_name == "功能场景和时序":
                for scenario_index, scenario in enumerate(normalized["scenarios"].values(), start=1):
                    _render_scenario(document, scenario, number=f"{section_number}.{scenario_index}")
                _render_signal_table(document, normalized["signals"], chapter.get("signal_keys", []))
            if chapter_id == "7":
                _render_traceability(document, normalized, section_name)

    document.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description="将 PRD 功能定义语义草稿渲染为不依赖模板的 DOCX")
    parser.add_argument("spec", help="function_definition_spec.json")
    parser.add_argument("--out", required=True, help="DOCX 输出目录")
    parser.add_argument("--name", help="DOCX 文件名")
    args = parser.parse_args()
    output = render_prd_spec_document(args.spec, args.out, args.name)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
