#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build reusable energy-management and charging-control DOCX templates.

The two source PDFs are signed project documents, not editable templates.  This
script intentionally creates a neutral skeleton: project names, controllers,
signals, thresholds and UI wording are left as placeholders for each task.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "templates"


def _paragraph(doc, text="", style=None, bold=False):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        run.bold = bold
    return p


def _table(doc, rows, cols):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    return table


def _set_field(run, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def _set_a4(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(20)


def _set_cell(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def _cover(doc):
    p = _paragraph(doc, "")
    p.paragraph_format.space_after = Pt(4)
    top = _table(doc, 1, 4)
    labels = ("交付物编号", "{{交付物编号}}", "交付物版本", "V{{文档版本}}")
    for cell, value in zip(top.rows[0].cells, labels):
        _set_cell(cell, value, bold=value in {"交付物编号", "交付物版本"})

    p = _paragraph(doc, "{{项目名称}}项目", bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(14)
    p = _paragraph(doc, "{{功能名称}}功能定义", bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(24)
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(30)

    approval = _table(doc, 9, 2)
    approval.autofit = True
    fields = (
        ("编制部门", "{{编制部门}}"),
        ("交付时间", "{{交付时间}}"),
        ("编    制", "{{编制人}}"),
        ("校    对", "{{校对人}}"),
        ("审    核", "{{审核人}}"),
        ("标准化", "{{标准化人}}"),
        ("审    定", "{{审定人}}"),
        ("会    签", "{{会签人}}"),
        ("批    准", "{{批准人}}"),
    )
    for row, (label, value) in zip(approval.rows, fields):
        _set_cell(row.cells[0], label, bold=True)
        _set_cell(row.cells[1], value)
    p = _paragraph(doc, "内部资料　注意保密")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].font.color.rgb = RGBColor(255, 0, 0)


def _cover_footer(section):
    section.different_first_page_header_footer = True
    footer = section.first_page_footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("页码 ")
    _set_field(p.add_run(), "PAGE")


def _body_header_footer(section):
    header = section.header
    table = header.add_table(rows=1, cols=3, width=Mm(165))
    table.style = "Table Grid"
    _set_cell(table.cell(0, 0), "{{公司名称}}", bold=True)
    _set_cell(table.cell(0, 1), "{{功能名称}}功能定义")
    _set_cell(table.cell(0, 2), "受控文件")
    table.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.cell(0, 2).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer
    table = footer.add_table(rows=2, cols=6, width=Mm(165))
    table.style = "Table Grid"
    row = table.rows[0].cells
    for cell, value in zip(row, ("编制", "{{编制人}}", "日期", "{{交付时间}}", "版本", "V{{文档版本}}")):
        _set_cell(cell, value, bold=value in {"编制", "日期", "版本"})
    row = table.rows[1].cells
    for cell, value in zip(row, ("文件名", "{{项目名称}}项目{{功能名称}}功能定义", "", "", "页码", "")):
        _set_cell(cell, value, bold=value == "文件名" or value == "页码")
    row[5].paragraphs[0].add_run(" ")
    _set_field(row[5].paragraphs[0].add_run(), "PAGE")
    row[5].paragraphs[0].add_run(" / ")
    _set_field(row[5].paragraphs[0].add_run(), "NUMPAGES")


def _change_record_and_toc(doc):
    _paragraph(doc, "更改记录表", bold=True)
    t = _table(doc, 2, 6)
    for cell, value in zip(t.rows[0].cells, ("序号", "生效日期", "更改人", "版本", "更改章节", "更改描述")):
        cell.text = value
    _paragraph(doc, "目录", bold=True)
    toc = (
        "1. 导言",
        "1.1 目的和适用范围",
        "1.2 参考文档",
        "1.3 措辞约定",
        "1.4 术语和缩略语",
        "2. 功能概述",
        "3. 系统功能描述",
        "3.1 主功能 Use Case",
        "3.2 子功能概述",
        "3.3 系统功能描述",
        "3.4 系统框图",
        "3.5 网络拓扑",
        "4. 系统需求",
        "4.1 功能需求",
        "4.2 车辆需求",
        "4.3 诊断需求",
        "4.4 FMEA 需求",
        "4.5 功能安全需求",
        "4.6 信息安全需求",
        "5. 功能定义",
        "5.x.1 概述",
        "5.x.2 框图",
        "5.x.3 接口",
        "5.x.4 变量/参数",
        "5.x.5 功能描述",
        "5.x.6 故障定义",
        "5.x.7 功能安全要求",
        "5.x.8 性能规范",
    )
    for item in toc:
        _paragraph(doc, item)


def _common(doc, domain_name):
    _paragraph(doc, "1. 导言", "Heading 1")
    _paragraph(doc, "1.1 目的和适用范围", "Heading 2")
    _paragraph(doc, "本文件定义{{功能名称}}的实现方案、系统边界、控制器职责和可验证需求，用于指导相关系统进行设计、开发、测试和验收。")
    _paragraph(doc, "本模板适用于{{项目名称}}项目中属于" + domain_name + "的功能；具体车型、平台、控制器、信号和标定值以项目输入及最新信号矩阵为准。")
    _paragraph(doc, "1.2 参考文档", "Heading 2")
    _paragraph(doc, "国家/行业/地方/企业标准、整车输入文件、供应商文件、对标报告及相关功能定义。")
    t = _table(doc, 2, 3)
    for cell, value in zip(t.rows[0].cells, ("序号", "名称", "版本")):
        cell.text = value
    _paragraph(doc, "1.3 措辞约定", "Heading 2")
    _paragraph(doc, "“必须”：无条件执行；“应该”：推荐执行；“建议”：可在评审确认后执行。正文中的规范性措辞应保持可追溯。")
    _paragraph(doc, "1.4 术语和缩略语", "Heading 2")
    t = _table(doc, 2, 3)
    for cell, value in zip(t.rows[0].cells, ("序号", "缩写/术语", "释义")):
        cell.text = value

    _paragraph(doc, "2. 功能概述", "Heading 1")
    _paragraph(doc, "列出本文件包含的主功能、功能 ID、功能边界和适用平台；不在此处展开逐条控制逻辑。")

    _paragraph(doc, "3. 系统功能描述", "Heading 1")
    _paragraph(doc, "3.1 主功能 Use Case", "Heading 2")
    _paragraph(doc, "按主功能填写目的和作用、参与者、前置条件、主成功场景、后置条件、备选/异常场景、特殊约束和涉及部件。")
    _paragraph(doc, "3.2 子功能概述", "Heading 2")
    _paragraph(doc, "用清单或表格建立主功能—子功能—主控/涉及控制器—引用关系，便于范围管理和需求追踪。")
    _paragraph(doc, "3.3 系统功能描述", "Heading 2")
    _paragraph(doc, "按 APP、T-BOX、ICC、PDCU/VDC、BMS、OBC/ECC 或项目实际控制器，描述职责、输入、输出、接口方向和跨控制器分工。")
    _paragraph(doc, "3.4 系统框图", "Heading 2")
    _paragraph(doc, "放置功能相关控制器、外部端、执行器和主要信息流的逻辑框图；关系不确定时标注待确认。")
    _paragraph(doc, "3.5 网络拓扑", "Heading 2")
    _paragraph(doc, "放置涉及 CAN/LIN/以太网/4G/蓝牙等网络的拓扑，标注网段、网关路由和功能相关节点。")

    _paragraph(doc, "4. 系统需求", "Heading 1")
    for title, desc in (
        ("4.1 功能需求", "按控制器或外部系统列出可验证的功能需求，并引用主功能/子功能 ID。"),
        ("4.2 车辆需求", "描述网关路由、网络管理、车辆状态和整车配置等跨功能约束。"),
        ("4.3 诊断需求", "定义通信丢失、输入无效、执行失败、超时和故障上报/清除条件。"),
        ("4.4 FMEA需求", "列出需纳入 FMEA 的失效模式、影响和预防/探测要求；无内容时明确暂无。"),
        ("4.5 功能安全需求", "引用适用的安全目标、降级策略、故障反应时间和安全状态；无适用项时明确暂无。"),
        ("4.6 信息安全需求", "描述远程控制、身份鉴权、通信保护、重放/伪造防护等要求；无适用项时明确暂无。"),
    ):
        _paragraph(doc, title, "Heading 2")
        _paragraph(doc, desc)


def _feature_module(doc, title, interface_sections):
    _paragraph(doc, "5. 功能定义", "Heading 1")
    _paragraph(doc, "功能模块", "Heading 2")
    _paragraph(doc, title, "Heading 3")
    for heading, desc in (
        ("概述", "说明功能目的、边界、主控制器和适用平台；避免写入只适用于某一项目的固定结论。"),
        ("框图", "放置本功能的控制流程或控制器交互框图。"),
        ("接口", "按下列接口类型分别列出信号/服务、方向、条件、周期/超时、范围、枚举和适用项目。"),
    ):
        _paragraph(doc, heading, "Heading 4")
        _paragraph(doc, desc)
    for section in interface_sections:
        _paragraph(doc, section, "Heading 4")
        _paragraph(doc, "填写本接口下的需求/信号表；没有该接口时删除本小节，不填“/”。")
    for heading, desc in (
        ("变量/参数", "列出标定量、阈值、计时器、默认值、范围、单位、精度及其来源；TBD 必须带责任人和关闭条件。"),
        ("功能描述", "按触发条件→判定/状态机→执行输出→反馈/显示→退出条件组织可验证逻辑，可拆分为多个子功能 ID。"),
        ("故障定义", "定义故障检测条件、故障等级、故障提示/上报、抑制条件和超时策略。"),
        ("功能安全要求", "填写与本功能直接相关的安全约束、降级、最小风险状态和验证要求；不适用时注明依据。"),
        ("性能规范", "填写时延、周期、成功率、持续时间、并发、功耗或温度等可测指标。"),
    ):
        _paragraph(doc, heading, "Heading 4")
        _paragraph(doc, desc)


def build(kind, display_name, domain_name, interface_sections):
    doc = Document()
    section = doc.sections[0]
    _set_a4(section)
    _cover(doc)
    _cover_footer(section)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    _set_a4(body_section)
    body_section.header.is_linked_to_previous = False
    body_section.footer.is_linked_to_previous = False
    _body_header_footer(body_section)
    _change_record_and_toc(doc)
    _common(doc, domain_name)
    _feature_module(doc, "功能模块 1", interface_sections)
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "宋体"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for style_name, size in (("Normal", 10.5), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11), ("Heading 4", 10.5)):
        if style_name in doc.styles:
            doc.styles[style_name].font.size = Pt(size)
    props = doc.core_properties
    props.title = display_name
    props.subject = "可复制功能章节模板"
    props.author = ""
    props.last_modified_by = ""
    out = OUT / kind / "template.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def main():
    build(
        "energy_management",
        "能量管理域功能定义模板",
        "能量管理/模式控制域",
        ("HMI", "硬线信号", "CAN 信号", "LIN 信号"),
    )
    build(
        "charging_control",
        "充电控制域功能定义模板",
        "充电控制/远程充电域",
        ("HMI", "硬线信号", "4G/蓝牙信号", "CAN 信号", "LIN 信号"),
    )


if __name__ == "__main__":
    main()
